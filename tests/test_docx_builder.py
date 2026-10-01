"""Open the produced .docx and assert RTL properties, merges, fonts, header/footer."""
import cv2
import numpy as np
import pytest
from docx import Document as open_docx
from docx.oxml.ns import qn

from img2docx.config import DocxCfg
from img2docx.docx_builder import build_docx
from img2docx.schema import BBox, Block, Document, Page, QAReport, SourceInfo, Style

TABLE = ('<table><tr><th rowspan="2">م</th><th rowspan="2">اسم المشروع</th><th colspan="2">التكلفة</th></tr>'
         '<tr><th>المعتمدة</th><th>المصروفة</th></tr>'
         '<tr><td>1</td><td>متحف المدينة</td><td>64000</td><td>1000</td></tr></table>')


def _doc(tmp_path, direction="rtl"):
    stamp = np.zeros((80, 80, 4), np.uint8)
    cv2.circle(stamp, (40, 40), 30, (160, 40, 40, 255), 4)
    cv2.imwrite(str(tmp_path / "stamp.png"), stamp)
    cv2.imwrite(str(tmp_path / "logo.png"), np.full((60, 60, 3), 80, np.uint8))
    bl = [
        Block(id="b1", type="header", order=0, bbox=BBox.of([1200, 100, 1550, 200]), text="وزارة الثقافة\nمديرية التراث", lang="ar", direction="rtl", style=Style(align="right", bold=True, size_pt=14)),
        Block(id="b2", type="logo", order=1, bbox=BBox.of([750, 80, 950, 280]), image_path="logo.png"),
        Block(id="b3", type="header", order=2, bbox=BBox.of([150, 100, 600, 200]), text="Ministry of Culture", lang="en", direction="ltr", style=Style(align="left")),
        Block(id="b4", type="doc_number", order=3, bbox=BBox.of([1300, 350, 1530, 400]), text="الرقم: ٤/٣٢٤١", lang="ar", direction="rtl", style=Style(align="right", size_pt=13)),
        Block(id="b5", type="paragraph", order=4, bbox=BBox.of([180, 500, 1530, 600]), text="يعقد مؤتمر Digital Heritage في القاعة الرئيسية.", lang="mixed", direction="rtl", style=Style(align="right", size_pt=13), confidence=0.5),
        Block(id="b6", type="table", order=5, bbox=BBox.of([180, 650, 1530, 900]), html=TABLE, lang="ar", direction="rtl"),
        Block(id="b7", type="paragraph", order=6, bbox=BBox.of([180, 950, 1530, 1000]), text="The plan has been approved.", lang="en", direction="ltr", style=Style(align="left")),
        Block(id="b8", type="stamp", order=7, bbox=BBox.of([600, 1000, 900, 1300]), image_path="stamp.png"),
        Block(id="b9", type="footer", order=8, bbox=BBox.of([650, 2250, 1000, 2290]), text="هاتف: 06 4600000", lang="ar", direction="rtl", style=Style(align="center")),
    ]
    return Document(source=SourceInfo(path="x.jpg", sha256="0", width_px=1697, height_px=2400),
                    page=Page(width_px=1697, height_px=2400, width_mm=210, height_mm=297, dpi=205, direction=direction),
                    blocks=bl, qa=QAReport(width_px=1697, height_px=2400, blur_laplacian_var=500, glare_ratio=0))


@pytest.fixture()
def built(tmp_path):
    doc = _doc(tmp_path)
    cfg = DocxCfg(highlight_low_confidence=True, comment_low_confidence=True)
    out = build_docx(doc, tmp_path / "t.docx", cfg, tmp_path)
    return open_docx(str(out)), cfg


def _runs_xml(p):
    return [r for r in p._p.iter(qn("w:r"))]


def test_section_is_rtl(built):
    d, _ = built
    assert d.sections[0]._sectPr.find(qn("w:bidi")) is not None
    assert abs(d.sections[0].page_width.mm - 210) < 0.5


def test_arabic_paragraph_bidi_and_rtl_runs(built):
    d, cfg = built
    p = next(p for p in d.paragraphs if "يعقد" in p.text)
    assert p._p.pPr.find(qn("w:bidi")) is not None and p._p.pPr.find(qn("w:bidi")).get(qn("w:val")) is None
    runs = _runs_xml(p)
    by_text = {"".join(t.text for t in r.iter(qn("w:t"))): r for r in runs}
    ar = next(r for t, r in by_text.items() if "يعقد" in t)
    en = next(r for t, r in by_text.items() if "Digital" in t)
    assert ar.rPr.find(qn("w:rtl")) is not None
    assert en.rPr.find(qn("w:rtl")) is None  # LTR run inside RTL paragraph
    rf = ar.rPr.find(qn("w:rFonts"))
    assert rf.get(qn("w:cs")) == cfg.font_cs and rf.get(qn("w:ascii")) == cfg.font_latin
    assert ar.rPr.find(qn("w:szCs")).get(qn("w:val")) == "26"  # 13pt complex-script size
    assert ar.rPr.find(qn("w:highlight")) is not None  # low confidence highlighted


def test_english_paragraph_is_ltr(built):
    d, _ = built
    p = next(p for p in d.paragraphs if "approved" in p.text)
    assert p._p.pPr.find(qn("w:bidi")).get(qn("w:val")) == "0"  # explicit: Normal style is bidi on RTL pages
    assert all(r.rPr.find(qn("w:rtl")) is None for r in _runs_xml(p))


def test_digits_preserved(built):
    d, _ = built
    assert any("٤/٣٢٤١" in p.text for p in d.paragraphs)


def test_table_rtl_and_merges(built):
    d, _ = built
    t = d.tables[0]
    assert t._tbl.tblPr.find(qn("w:bidiVisual")) is not None
    xml = t._tbl.xml
    assert 'w:gridSpan w:val="2"' in xml
    assert xml.count("<w:vMerge") >= 2  # restart + continue
    assert len(t.columns) == 4 and len(t.rows) == 3
    assert t.cell(0, 0).text == "م" and t.cell(1, 0).text == "م"  # rowspan cell spans rows 0-1
    assert t.cell(0, 2).text == "التكلفة" and t.cell(0, 3).text == "التكلفة"


def test_header_footer_and_images(built):
    d, _ = built
    hdr = d.sections[0].header
    htxt = "\n".join(c.text for t in hdr.tables for row in t.rows for c in row.cells)
    assert "وزارة الثقافة" in htxt and "Ministry of Culture" in htxt
    ht = hdr.tables[0]
    assert ht._tbl.tblPr.find(qn("w:bidiVisual")) is not None
    assert "وزارة" in ht.cell(0, 0).text  # visual right cell (bidiVisual) holds the Arabic letterhead
    assert "Ministry" in ht.cell(0, 2).text
    assert hdr._element.xml.count("<pic:pic") == 1  # logo
    ftxt = "\n".join(p.text for p in d.sections[0].footer.paragraphs)
    assert "هاتف" in ftxt
    body = d.element.body.xml
    assert "<wp:anchor" in body and 'behindDoc="1"' in body  # floating stamp


def test_comment_added_for_low_confidence(built):
    d, _ = built
    assert len(list(d.comments)) >= 1


def test_ltr_page(tmp_path):
    doc = _doc(tmp_path, direction="ltr")
    out = build_docx(doc, tmp_path / "l.docx", DocxCfg(), tmp_path)
    d = open_docx(str(out))
    assert d.sections[0]._sectPr.find(qn("w:bidi")) is None
